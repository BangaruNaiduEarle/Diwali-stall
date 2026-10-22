#!/usr/bin/env py -3
"""Read-only dump of CRACKERS invoice line items (both sheets). Stdin JSON, stdout JSON."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

HEADER_ROW = 4
DATA_START = 5


def cell_value(ws, row: int, col: int):
    return ws.cell(row, col).value


def is_line_row(sl, particulars) -> bool:
    if particulars is None or str(particulars).strip() == "":
        return False
    text = str(particulars).strip().lower()
    if text in {"particulars", "total", "grand total"}:
        return False
    if sl is None:
        return True
    sl_text = str(sl).strip().lower()
    return sl_text not in {"sl. no", "sl no", "s.no"}


def read_workbook(workbook_path: str) -> dict:
    path = Path(workbook_path)
    if not path.is_file():
        return {"ok": False, "error": f"workbook_not_found:{path.name}", "lines": []}

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    lines = []
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            row_num = DATA_START
            for row in ws.iter_rows(min_row=DATA_START, max_col=8, values_only=True):
                vals = list(row) + [None] * (8 - len(row))
                sl, comp, particulars, quantity, per, rate, x4, amount = vals[:8]
                if is_line_row(sl, particulars):
                    lines.append(
                        {
                            "workbook": path.name,
                            "sheet": sheet_name,
                            "excel_row": row_num,
                            "sl": sl,
                            "comp": None if comp is None else str(comp).strip(),
                            "particulars": str(particulars).strip(),
                            "quantity": quantity,
                            "per": None if per is None else str(per).strip(),
                            "rate": rate,
                            "x4": x4,
                            "amount": amount,
                        }
                    )
                row_num += 1
    finally:
        wb.close()
    return {"ok": True, "workbook": path.name, "lines": lines}


def main() -> int:
    try:
        payload = json.load(sys.stdin)
    except json.JSONDecodeError as exc:
        print(json.dumps({"ok": False, "error": f"invalid_stdin_json:{exc}", "lines": []}))
        return 1
    workbook_path = str(payload.get("workbook_path") or "")
    out = read_workbook(workbook_path)
    json.dump(out, sys.stdout, default=str)
    return 0 if out.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
