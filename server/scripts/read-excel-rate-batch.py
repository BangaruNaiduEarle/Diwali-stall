#!/usr/bin/env py -3
"""Read-only batch read of Excel bill rows (Rate col 6, X4 col 7). Stdin JSON, stdout JSON."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


def cell_value(ws, row: int, col: int):
    return ws.cell(row, col).value


def read_batch(payload: dict) -> dict:
    cells = payload.get("cells") or []
    by_workbook: dict[str, list[dict]] = {}
    for entry in cells:
        wb_path = str(entry.get("workbook_path") or "")
        by_workbook.setdefault(wb_path, []).append(entry)

    results = []
    for wb_path, group in by_workbook.items():
        path = Path(wb_path)
        if not path.is_file():
            for entry in group:
                results.append(
                    {
                        "key": entry.get("key"),
                        "error": f"workbook_not_found:{path.name}",
                    }
                )
            continue
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet_cache: dict[str, object] = {}
        try:
            for entry in group:
                key = entry.get("key")
                sheet_name = str(entry.get("worksheet") or "")
                row = int(entry.get("row") or 0)
                if row < 1:
                    results.append({"key": key, "error": "invalid_row"})
                    continue
                if sheet_name not in wb.sheetnames:
                    results.append({"key": key, "error": f"worksheet_not_found:{sheet_name}"})
                    continue
                if sheet_name not in sheet_cache:
                    sheet_cache[sheet_name] = wb[sheet_name]
                ws = sheet_cache[sheet_name]
                try:
                    rate = cell_value(ws, row, 6)
                    x4 = cell_value(ws, row, 7)
                    results.append(
                        {
                            "key": key,
                            "rate_raw": rate,
                            "x4_raw": x4,
                            "error": None,
                        }
                    )
                except Exception as exc:  # noqa: BLE001
                    results.append({"key": key, "error": f"read_error:{exc}"})
        finally:
            wb.close()
    return {"results": results}


def main() -> int:
    try:
        payload = json.load(sys.stdin)
    except json.JSONDecodeError as exc:
        print(json.dumps({"error": f"invalid_stdin_json:{exc}"}))
        return 1
    out = read_batch(payload)
    json.dump(out, sys.stdout)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
